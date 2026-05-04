import sys
import glob
import os
import re
import pandas as pd
import time
import warnings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter

# Конфигурационные параметры
# тут инфекци без дублей (иначе дубль в рез-тах)
CATEGORIES = {
            'Заяц РНК': ['FCV', 'CPIV', 'FCoV', 'CCoV', 'CDV', 'Нью'],
            'Заяц ДНК': ['Giard', 'Salm', 'Tritri', 'Bartonella', 'Токсопл', 'Microsp','Trich', 'gibsoni', 'Орнитоз', 'Полио', 'Цирко'],
            'Genlab': ['HV', 'Mycoplasma spp', 'M.felis', 'Chlamyd', 'M.canis', 'FIV', 'FeLV', 'Campilobacter jejuni', 'Клостр', 'FIV+FeLV'],
            'Fractal': ['Борд', 'PV', 'Crypto', 'CAV', 'Anaplasma', 'Борр', 'Ehr', 'Urea', 'Диро', 'Lepto'],
            'VectBest': ['Bruc','haemofelis', 'haеmocanis', 'perfringens', 'galiseptica', 'Паст','Асперг', 'Babesia canis','Babesia spp', 'РНК FeLV', 'Haemobartonella felis', 'Haemobartonella canis', 'Campylobacter spp']
        }


PRIORITY_COMPLEXES = {
    'ПЦР-РЕСП-К': ['FCV', 'HV', 'Mycoplasma spp', 'M.felis','Борд', 'Chlamyd'],
    'ПЦР-РЕСП-С': ['CPIV','HV', 'Mycoplasma spp','Борд', 'M.canis', 'CAV'],
    'ПЦР-ДИАР-К': ['FCoV', 'Giard', 'Salm','PV','Campilobacter jejuni', 'Клостр','Crypto'],
    'ПЦР-ДИАР-С': ['CCoV', 'СCoV','Giard', 'Salm','PV','Campilobacter jejuni', 'Клостр', 'Crypto']
}

def normalize_keys(key):
            key = key.replace('М', 'M') 
            key = key.replace('С', 'C')
            return key

def process_data(df):
    ##Обработка всех данных и только вывод сначала инфекций из комплексов, а потом обычных      
    # тут ВСЕ ВАРИАНТЫ написания
    infections = ['FCV', 'CPIV', 'FCoV', 'CCoV', 'СCoV', 'CDV', 'Нью', 'Giard', 'Salm', 'Tritri', 'Bartonella', 'Токсопл', 'gibsoni',
                      'Асперг', 'Орнитоз', 'Полио', 'Цирко','Bruc', 'haemofelis', 'haеmocanis', 'perfringens', 'galiseptica', 'Паст', 'B.canis', 'Babesia canis',
                      'B.spp', 'Babesia spp', 'РНК', 'HV', 'Мycoplasma spp', 'Mycoplasma spp', 'M.felis', 'Mycoplasma felis','FIV', 'FeLV', 'PV', 'Campylobacter spp', 'Campilobacter jejuni', 'Клостр', 'Борд', 'Chlamyd', 
                      'Crypto', 'CAV', 'M.canis', 'Anaplasma', 'Борр', 'Ehr', 'Urea', 'Диро', 'Lepto', 'Microsp', 'Trich', 'Bruc',
                      'Haemobartonella felis', 'Haemobartonella canis', 'Бабезиоз собак и кошек']
    

    results = {} 
    bcanis_nums1 = []
    bcanis_nums2 = []
    gibsoni_nums  = []
    bspp_numbers1  = []
    bspp_numbers2  = []
    bspp_imposters = []
    felis_nums = []
    felis_new = []

    for infection in infections:
        result_key = normalize_keys(infection.strip())
        if (result_key == "Mycoplasma spp"):
            pattern = r'[MМ]ycoplasma\s+spp'
            filtered = df[df.iloc[:, 4].str.contains(pattern, case=True, na=False, regex=True)]
        elif (result_key == "CCoV"):
            pattern = r'[CС]oV'
            mask_c = df.iloc[:, 4].str.contains(pattern, case=True, na=False, regex=True)
            mask_f = df.iloc[:, 4].str.contains('FCoV', case=True, na=False, regex=False)
            filtered = df[mask_c & ~mask_f]
        elif (result_key == "FCoV"):
            mask_f = df.iloc[:, 4].str.contains('FCoV', case=True, na=False, regex=False)
            filtered = df[mask_f]
        else:
            filtered = df[df.iloc[:, 4].str.contains(result_key, case=True, na=False, regex=False)]

        results[result_key] = filtered.iloc[:, 0].unique().tolist()
                                  
    if 'Бабезиоз собак и кошек' in results:
        bspp_imposters = results.get('Бабезиоз собак и кошек', [])
    if 'B.canis' in results:
        bcanis_nums1 = results.get('B.canis', [])
    if 'Babesia canis' in results:
        bcanis_nums2 = results.get('Babesia canis', [])
    if 'Babesia spp' in results:
        bspp_numbers1 = results.get('Babesia spp', [])
    
    if 'B.spp' in results:
        bspp_numbers2 = results.get('B.spp', [])
    if 'gibsoni' in results:
        gibsoni_nums = results.get('gibsoni', [])

    if 'M.felis' in results:
        felis_nums = results.get('M.felis', [])
    if 'Mycoplasma felis' in results:
        felis_new = results.get('Mycoplasma felis', [])

    if felis_new:
        felis_nums = set(felis_nums + felis_new) 
        
    results['M.felis'] = [num for num in felis_nums]
    results['M.felis'] = sorted(results['M.felis'] )
    
    bcanis_nums = set(bcanis_nums1  + bcanis_nums2)
    bspp_numbers = set(bspp_numbers1 + bspp_numbers2 + bspp_imposters) 


    results['Babesia canis'] = [num for num in bcanis_nums]
    results['gibsoni'] = [num for num in gibsoni_nums]
    results['Babesia spp'] = [num for num in bspp_numbers]
    
    if bspp_imposters:
        #results.setdefault('Babesia spp', []).extend(bspp_numbers)
        results['Babesia spp'] = [num for num in bspp_numbers]
        results['Babesia spp'] = sorted(results['Babesia spp'])
        results['Babesia canis'] = [num for num in bcanis_nums if num not in bspp_imposters]
        results['Babesia canis'] = sorted(results['Babesia canis'])
        results['gibsoni'] = [num for num in gibsoni_nums if num not in bspp_imposters]
        results['gibsoni'] = sorted(results['gibsoni'])
        
    #FIV+FeLV
    fiv = set(results.get('FIV', []))
    felv = set(results.get('FeLV', []))
    results['FIV+FeLV'] = sorted(fiv & felv)
    results['FIV'] = sorted(fiv - felv)
    results['FeLV'] = sorted(felv - fiv)
    if 'РНК' in results:
        results['РНК FeLV'] = results.pop('РНК')

    for complex_name, inf_list in PRIORITY_COMPLEXES.items():
        #mask = (df.iloc[:, 4].str.contains(complex_name, na=False) & ~df.iloc[:, 4].str.contains(complex_name+'Б1', na=False) & ~df.iloc[:, 4].str.contains(complex_name+'Б2', na=False))
        #filtered_complexes = df[mask]
        escaped = re.escape(complex_name)
        pattern = f"{escaped}(?![\\w\\d])"
        mask = df.iloc[:, 4].str.contains(pattern, na=False, regex=True)
        filtered_complexes = df[mask]
        results[complex_name] = filtered_complexes.iloc[:, 0].tolist()
    #print(results)
    return results

def create_excel_report(data):
    #Создание отчета с разделением на приоритетные и обычные блоки
    wb = Workbook()
    ws = wb.active
    
    # Стили офддормления
    thin_border = Border(
        left=Side(style='thin'), 
        right=Side(style='thin'),
        top=Side(style='thin'), 
        bottom=Side(style='thin')
    )
    thick_border = Border(
        left=Side(style='medium'), 
        right=Side(style='medium'),
        top=Side(style='medium'), 
        bottom=Side(style='medium')
    )
    category_font = Font(bold=True, size=14)
    header_font = Font(bold=True, size=12)
    
    current_row = 1
    

    ####
    def write_block(category, infections, block_type, data_dict):
        #Запись блока данных для категории
        nonlocal current_row
        cols = []
        for inf in infections:
            chunks = 1
            while f"{inf}_{chunks}" in data_dict:
                cols.append(f"{inf}_{chunks}")
                chunks += 1

        if not cols:
            return
    ######
        
        #Заголовок блока
        ws.merge_cells(
            start_row=current_row,
            end_row=current_row,
            start_column=1,
            end_column=len(cols)
        )
        cell = ws.cell(current_row, 1, f"{category} ({block_type})")
        cell.font = category_font
        cell.alignment = Alignment(horizontal='center')
        #Границы заголовка
        for col in range(1, len(cols)+1):
            ws.cell(current_row, col).border = thick_border
        
        current_row += 1
        
        #Заголовки столбцов
        for col_idx, col in enumerate(cols, 1):
            cell = ws.cell(current_row, col_idx, col.split('_')[0])
            cell.font = header_font
            cell.border = thin_border
        
        current_row += 1
        
        #Данные
        max_rows = max(len(data_dict[col]) for col in cols)
        for i in range(max_rows):
            for col_idx, col in enumerate(cols, 1):
                val = data_dict[col][i] if i < len(data_dict[col]) else ''
                ws.cell(current_row, col_idx, val).border = thin_border
            current_row += 1
        
        current_row += 1
    
    #Форматирование данных
    def format_data(data):
        formatted = {}
        for infection, numbers in data.items():
            chunks = [numbers[i:i+8] for i in range(0, len(numbers), 8)]
            for i, chunk in enumerate(chunks, 1):
                formatted[f"{infection}_{i}"] = chunk + ['']*(8-len(chunk))
        return formatted
    

    #print(f"DATA BEFORE: {data}")
     # Собираем все приоритетные инфекции из PRIORITY_COMPLEXES
    complexes_dict = {}
    for complex_name, inf_list in PRIORITY_COMPLEXES.items():
        if complex_name in data:
            complex_numbers = data[complex_name]
            #print(complex_name)
            for inf in inf_list:
                if inf not in complexes_dict:
                    complexes_dict[inf] = []
                complexes_dict[inf].extend(complex_numbers)
            del data[complex_name]

    #print(f"DATA AFTER: {data}") ##OK
    #print(complexes_dict)


    other_data_set = set()
    complexes_set = {(key, int(value)) for key in complexes_dict for value in complexes_dict[key]}
    all_data_set = {(key, int(value)) for key in data for value in data[key]}
    other_data_set = all_data_set - complexes_set
    #print(complexes_set)

    ##
    def func_set_to_dict(cur_set):
        out_data = {}
        for key, value in cur_set:
            if key not in out_data:
                out_data[key] = []
            out_data[key].append(value)

        for key in out_data:
            out_data[key].sort()
        
        return out_data
    ###

    complexes_data = func_set_to_dict(complexes_set)
    other_data = func_set_to_dict(other_data_set)
    #print(complexes_data)
    #print(other_data)

    formatted_complex_data = format_data(complexes_data)
    formatted_other_data = format_data(other_data)

    # Обрабатываем каждую категорию
    for category, infections in CATEGORIES.items():
        #Данные комплексов
        write_block(category, infections, 'Комплексы', formatted_complex_data)

        #Данные остальные
        write_block(category, infections, 'Не-комплексы', formatted_other_data)
    
    #Автонастройка ширины
    for col in ws.columns:
        max_len = 0
        for cell in col:
            try:
                if len(str(cell.value)) > max_len:
                    max_len = len(str(cell.value))
            except:
                pass
        if max_len > 0:
            ws.column_dimensions[get_column_letter(col[0].column)].width = (max_len + 2) * 1.2
    
    return wb

def main():
    try:
        input_files = glob.glob('*.xlsx')
        if len(input_files) != 1:
            raise ValueError(f"Найдено {len(input_files)} файлов. Требуется 1.")
        
        input_path = input_files[0]
        output_dir = os.path.abspath(os.path.join(os.getcwd(), os.pardir))
        time_suffix = "ночь" if int(time.strftime("%H")) > 12 else "день"
        output_path = os.path.join(
            output_dir,
            f"postanovka_{time.strftime('%d%m%y')}_{time_suffix}_v2.xlsx"
        )
        
        df = pd.read_excel(input_path, sheet_name=0)
        df.iloc[:,4] = df.iloc[:,4].astype(str).str.strip()
        
        df.iloc[:, 0] = df.iloc[:, 0].ffill()
        processed_data = process_data(df)
        
        report = create_excel_report(processed_data)
        report.save(output_path)
        
        print(f"Результаты записаны в файл: {output_path}")
    
    except Exception as e:
        print(f"Error!{e}")
        sys.exit(1)

if __name__ == '__main__':
    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl.styles.stylesheet")
    main()