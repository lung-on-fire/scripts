import sys
import glob
import os
import pandas as pd
import time
import warnings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter

def read_file():
    try:
        files = glob.glob('*.xlsx')
        if len(files) != 1:
            raise ValueError(f"Найдено {len(files)} файлов. Нужен ровно один XLSX файл.")
        
        input_file = files[0]
        parent_dir = os.path.abspath(os.path.join(os.getcwd(), os.pardir))
        time_suffix = "ночь_v1" if int(time.strftime("%H")) > 12 else "день_v1"
        output_file = os.path.join(parent_dir, f"postanovka_{time.strftime('%d%m%y')}_{time_suffix}.xlsx")

        # Стили оформления
        header_font = Font(bold=True, size=12)
        category_font = Font(bold=True, size=14)
        thick_border = Border(
            left=Side(style='medium'),
            right=Side(style='medium'),
            top=Side(style='medium'),
            bottom=Side(style='medium')
        )

        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        data = pd.read_excel(input_file, sheet_name=0)
        data.iloc[:, 4] = data.iloc[:, 4].astype(str).str.strip()

        #Заполнение первого столбика сплошь номерами
        data.iloc[:, 0] = data.iloc[:, 0].ffill()

        # Сбор и обработка данных
        infections = ['FCV', 'CPIV', 'FCoV', 'CCoV', 'СCoV', 'CDV', 'Нью', 'Giard', 'Salm', 'Tritri', 'Bartonella', 'Токсопл', 'gibsoni',
                      'Асперг', 'Орнитоз', 'Полио', 'Цирко','Bruc', 'haemofelis', 'haеmocanis', 'perfringens', 'galiseptica', 'Паст', 'B.canis', 'Babesia canis',
                      'B.spp', 'Babesia spp', 'РНК', 'HV', 'Мycoplasma spp', 'Mycoplasma spp', 'M.felis', 'Mycoplasma felis', 'FIV', 'FeLV', 'PV', 'Campylobacter spp', 'Campilobacter jejuni', 'Клостр', 'Борд', 'Chlamyd', 
                      'Crypto', 'CAV', 'M.canis', 'Anaplasma', 'Борр', 'Ehr', 'Urea', 'Диро', 'Lepto', 'Microsp', 'Trich', 'Bruc',
                      'Haemobartonella felis', 'Haemobartonella canis', 'Бабезиоз собак и кошек']


        results = {}
        bcanis_nums1 = []
        bcanis_nums2 = []
        gibsoni_nums  = []
        bspp_numbers1  = []
        bspp_numbers2  = []
        bspp_imposters = []
        felis_nums = []
        felis_new = []

        def normalize_keys(key):
            key = key.replace('М', 'M') 
            key = key.replace('С', 'C')

            return key
        
        for infection in infections:
            result_key = normalize_keys(infection.strip())

            if (result_key == "Mycoplasma spp"):
                pattern = r'[MМ]ycoplasma\s+spp'
                infection_data = data[data.iloc[:, 4].str.contains(pattern, case=True, na=False, regex=True)]

            elif (result_key == "CCoV"):
                pattern = r'[CС]oV'
                mask_c = data.iloc[:, 4].str.contains(pattern, case=True, na=False, regex=True)
                mask_f = data.iloc[:, 4].str.contains('FCoV', case=True, na=False, regex=False)
                infection_data = data[mask_c & ~mask_f]
            
            elif (result_key == "FCoV"):
                mask_f = data.iloc[:, 4].str.contains('FCoV', case=True, na=False, regex=False)
                infection_data = data[mask_f]

            else:
                infection_data = data[data.iloc[:, 4].str.contains(result_key, case=True, na=False, regex=False)]
            
            infection_numbers = infection_data.iloc[:, 0]
            results[result_key] = [int(p) for p in infection_numbers.tolist()]

        if 'Бабезиоз собак и кошек' in results:
            bspp_imposters = results.get('Бабезиоз собак и кошек', [])

        if 'B.canis' in results:
            bcanis_nums1 = results.get('B.canis', [])
        if 'Babesia canis' in results:
            bcanis_nums2 = results.get('Babesia canis', [])

        if 'Babesia spp' in results:
            bspp_numbers1 = results.get('Babesia spp', [])
        
        if 'B.spp' in results:
            bspp_numbers2 = results.get('B.spp', [])

        if 'gibsoni' in results:
            gibsoni_nums = results.get('gibsoni', [])

        if 'M.felis' in results:
            felis_nums = results.get('M.felis', [])
        if 'Mycoplasma felis' in results:
            felis_new = results.get('Mycoplasma felis', [])

        if felis_new:
            felis_nums = set(felis_nums + felis_new) 
        
        results['M.felis'] = [num for num in felis_nums]
        results['M.felis'] = sorted(results['M.felis'] )

        bcanis_nums = set(bcanis_nums1  + bcanis_nums2)
        bspp_numbers = set(bspp_numbers1 + bspp_numbers2 + bspp_imposters) 

        results['Babesia canis'] = [num for num in bcanis_nums]
        results['gibsoni'] = [num for num in gibsoni_nums]
        results['Babesia spp'] = [num for num in bspp_numbers]

        if bspp_imposters:
            #results.setdefault('Babesia spp', []).extend(bspp_numbers)
            results['Babesia spp'] = [num for num in bspp_numbers]
            results['Babesia spp'] = sorted(results['Babesia spp'])

            results['Babesia canis'] = [num for num in bcanis_nums if num not in bspp_imposters]
            results['Babesia canis'] = sorted(results['Babesia canis'])

            results['gibsoni'] = [num for num in gibsoni_nums if num not in bspp_imposters]
            results['gibsoni'] = sorted(results['gibsoni'])


            
        #FIV+FeLV
        fiv = set(results.get('FIV', []))
        felv = set(results.get('FeLV', []))
        results['FIV+FeLV'] = sorted(fiv & felv)
        results['FIV'] = sorted(fiv - felv)
        results['FeLV'] = sorted(felv - fiv)
        if 'РНК' in results:
            results['РНК FeLV'] = results.pop('РНК')

        # Подготовка данных для транспонированной записи
        formatted_data = {}
        for infection, numbers in results.items():
            # Разбиваем на чанки по 8 номеров
            chunks = [numbers[i:i+8] for i in range(0, len(numbers), 8)]
            for i, chunk in enumerate(chunks):
                # Дополняем каждый чанк до 8 элементов пустыми строками
                chunk += [''] * (8 - len(chunk))
                formatted_data[f"{infection}_{i+1}"] = chunk

        # Создаем DataFrame и транспонируем его
        results_df = pd.DataFrame(formatted_data)
        results_df = results_df.T

        categories = {
            'Заяц РНК': ['FCV', 'CPIV', 'FCoV', 'CCoV', 'CDV', 'Нью'],
            'Заяц ДНК': ['Giard', 'Salm', 'Tritri', 'Bartonella', 'Токсопл', 'Microsp','Trich', 'gibsoni','Орнитоз', 'Полио', 'Цирко'],
            'Genlab': ['HV', 'Mycoplasma spp', 'M.felis', 'Chlamyd', 'M.canis', 'FIV', 'FeLV', 'Campilobacter jejuni', 'Клостр', 'FIV+FeLV'],
            'Fractal': ['Борд', 'PV', 'Crypto', 'CAV', 'Anaplasma', 'Борр', 'Ehr', 'Urea', 'Диро', 'Lepto'],
            'VectBest': ['Bruc','haemofelis', 'haеmocanis', 'perfringens', 'galiseptica', 'Паст','Асперг','Babesia canis', 'Babesia spp', 'РНК FeLV', 'Haemobartonella felis', 'Haemobartonella canis', 'Campylobacter spp']
        }

        wb = Workbook()
        ws = wb.active
        ws.title = "Результаты"
        current_row = 1

        for cat_name, infections in categories.items():
            # Фильтруем строки для текущей категории
            rows_to_include = [row_name for row_name in results_df.index if any(inf in row_name for inf in infections)]
            if not rows_to_include:
                continue

            # Разбиваем строки на блоки по 12 строк
            blocks = []
            for i in range(0, len(rows_to_include), 12):
                block = rows_to_include[i:i+12]
                blocks.append(block)

            # Обрабатываем каждый блок
            for block_idx, block_rows in enumerate(blocks):
                if current_row > 1:
                    current_row += 2  # Добавляем отступ между блоками

                # Заголовок категории (с номером блока если нужно)
                category_header = cat_name
                if len(blocks) > 1:
                    category_header = f"{cat_name} (блок {block_idx + 1})"
                
                ws.merge_cells(
                    start_row=current_row,
                    end_row=current_row,
                    start_column=1,
                    end_column=9  # 8 колонок данных + 1 колонка для названий анализов
                )
                cell = ws.cell(row=current_row, column=1, value=category_header)
                cell.font = category_font
                cell.alignment = Alignment(horizontal='center')
                
                for col in range(1, 10):
                    ws.cell(row=current_row, column=col).border = thick_border
                
                current_row += 1

                # Записываем данные для текущего блока
                for row_name in block_rows:
                    # Название анализа (первая колонка)
                    cell = ws.cell(row=current_row, column=1, value=row_name)
                    cell.font = header_font
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal='left')
                    
                    # Данные (8 колонок)
                    row_data = results_df.loc[row_name]
                    for col_idx in range(8):  # 8 колонок данных
                        value = row_data.iloc[col_idx] if col_idx < len(row_data) else ''
                        cell = ws.cell(row=current_row, column=col_idx + 2, value=value)
                        cell.border = thin_border
                        cell.alignment = Alignment(horizontal='center')
                    
                    current_row += 1

                # Заполняем оставшиеся строки пустыми значениями до 12 строк в блоке
                while len(block_rows) < 12:
                    for col in range(1, 10):  # 9 колонок (1 название + 8 данных)
                        cell = ws.cell(row=current_row, column=col, value='')
                        cell.border = thin_border
                    current_row += 1
                    block_rows.append('')

        # Автонастройка ширины столбцов
        for col in ws.columns:
            max_length = 0
            column_number = col[0].column
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min((max_length + 2) * 1.2, 25)  # Ограничиваем максимальную ширину
            col_letter = get_column_letter(column_number)
            ws.column_dimensions[col_letter].width = adjusted_width

        wb.save(output_file)
        print(f"Результаты записаны в файл: {output_file}")

    except Exception as e:
        print(f"Error! {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)

if __name__ == '__main__':
    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl.styles.stylesheet")
    read_file()